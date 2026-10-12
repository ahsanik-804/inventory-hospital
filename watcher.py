"""
Clinic Inventory — Excel Watcher + Supabase Sync
Watches inventory.xlsx → syncs to Supabase in real-time
"""

import json, sys, time, threading
from pathlib import Path

SUPABASE_URL = "https://hnekjyxobrezwkekrwnb.supabase.co"
SUPABASE_KEY = "eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJpc3MiOiJzdXBhYmFzZSIsInJlZiI6ImhuZWtqeXhvYnJlendrZWtyd25iIiwicm9sZSI6ImFub24iLCJpYXQiOjE3ODMwNTk0MTMsImV4cCI6MjA5ODYzNTQxM30.-gjskSdaDhkZEt7NANm7FDzCFD0ntMJKJ3bKppYa6zo"
EXCEL_FILENAME = "inventory.xlsx"

BASE_DIR   = Path(__file__).parent.resolve()
EXCEL_PATH = BASE_DIR / EXCEL_FILENAME

try:
    import openpyxl
except ImportError:
    print("Installing openpyxl..."); import subprocess
    subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl", "--quiet"])
    import openpyxl

try:
    from watchdog.observers import Observer
    from watchdog.events import FileSystemEventHandler
except ImportError:
    print("Installing watchdog..."); import subprocess
    subprocess.run([sys.executable, "-m", "pip", "install", "watchdog", "--quiet"])
    from watchdog.observers import Observer
    from watchdog.events import FileSystemEventHandler

try:
    import requests
except ImportError:
    print("Installing requests..."); import subprocess
    subprocess.run([sys.executable, "-m", "pip", "install", "requests", "--quiet"])
    import requests


def excel_to_rows(excel_path):
    try:
        wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    except Exception as e:
        print(f"  [!] Cannot read Excel: {e}"); return None

    all_rows = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2: continue

        headers = {str(h).strip().lower(): idx for idx, h in enumerate(rows[0]) if h is not None}

        def get(row, *names):
            for n in names:
                idx = headers.get(n.lower())
                if idx is not None and idx < len(row):
                    v = row[idx]
                    return v if v is not None else ""
            return ""

        for row in rows[1:]:
            name    = str(get(row, "name", "medicinename")).strip()
            company = str(get(row, "company")).strip()
            if not name or not company: continue

            try:    price = float(get(row, "price") or 0)
            except: price = 0
            try:    stock = int(float(get(row, "stock") or 0))
            except: stock = 0

            all_rows.append({
                "category":     sheet_name.strip(),
                "company":      company,
                "disease":      str(get(row, "disease")).strip(),
                "medicine_name": name,
                "salt":         str(get(row, "salt", "formula")).strip(),
                "weight":       str(get(row, "weight")).strip(),
                "price":        price,
                "stock":        stock
            })
    wb.close()
    return all_rows


def sync_to_supabase(rows):
    headers = {
        "apikey": SUPABASE_KEY,
        "Authorization": f"Bearer {SUPABASE_KEY}",
        "Content-Type": "application/json",
        "Prefer": "resolution=merge-duplicates"
    }

    # Delete all existing rows first
    del_res = requests.delete(
        f"{SUPABASE_URL}/rest/v1/medicines?id=gt.0",
        headers=headers
    )

    if del_res.status_code not in (200, 204):
        print(f"  [!] Delete failed: {del_res.status_code} {del_res.text}")
        return False

    if not rows:
        print("  [OK] Supabase cleared — no rows in Excel")
        return True

    # Insert all rows in batches of 100
    batch_size = 100
    for i in range(0, len(rows), batch_size):
        batch = rows[i:i+batch_size]
        ins_res = requests.post(
            f"{SUPABASE_URL}/rest/v1/medicines",
            headers=headers,
            json=batch
        )
        if ins_res.status_code not in (200, 201):
            print(f"  [!] Insert failed: {ins_res.status_code} {ins_res.text}")
            return False

    print(f"  [OK] Supabase synced — {len(rows)} medicines")
    return True


def refresh(reason=""):
    if not EXCEL_PATH.exists():
        print(f"  [!] File not found: {EXCEL_PATH}")
        return

    rows = excel_to_rows(EXCEL_PATH)
    if rows is None: return

    tag = f" [{reason}]" if reason else ""
    print(f"  Syncing {len(rows)} medicines to Supabase{tag}...")
    sync_to_supabase(rows)


_last_event = 0
_debounce   = 2.0

class ExcelHandler(FileSystemEventHandler):
    def on_modified(self, event):
        global _last_event
        if Path(event.src_path).resolve() != EXCEL_PATH: return
        now = time.time()
        if now - _last_event < _debounce: return
        _last_event = now
        threading.Timer(1.0, refresh, kwargs={"reason": "Excel saved"}).start()

    def on_created(self, event):
        self.on_modified(event)


if __name__ == "__main__":
    print("=" * 54)
    print("  Clinic Inventory — Supabase Watcher")
    print("=" * 54)
    print(f"  Excel  : {EXCEL_PATH}")
    print(f"  Cloud  : {SUPABASE_URL}")
    print("-" * 54)
    print("  Initial sync on startup...")
    refresh("startup")

    observer = Observer()
    observer.schedule(ExcelHandler(), str(BASE_DIR), recursive=False)
    observer.start()

    print(f"\n  Watching: {BASE_DIR}")
    print("  Save Excel → Supabase auto-updates → All devices sync.")
    print("  Press Ctrl+C to stop.")
    print("-" * 54)

    try:
        while True: time.sleep(1)
    except KeyboardInterrupt:
        observer.stop()
        print("\n  Watcher stopped.")
    observer.join()